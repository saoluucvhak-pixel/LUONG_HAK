from __future__ import annotations
import csv
from datetime import datetime
import json
from pathlib import Path
from .attendance_audit import audit_attendance
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def safe_value(text: str):
    """Guard exported CSV/XLSX from spreadsheet formula injection."""
    if isinstance(text,str) and text.lstrip().startswith(('=','+','-','@')):
        return "'" + text
    return text


def result_to_json(result: dict, destination: Path):
    destination.parent.mkdir(parents=True,exist_ok=True)
    destination.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return destination


def result_to_csv(result: dict, destination: Path):
    destination.parent.mkdir(parents=True,exist_ok=True)
    with destination.open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.writer(f)
        writer.writerow(['Trường','Nội dung','Đã duyệt','Trang'])
        for field in result.get('fields',[]):
            writer.writerow([safe_value(field['label']),safe_value(str(field['value'])),int(field['verified']),field['page']+1])
        for page in result.get('pages',[]):
            for t in page.get('tables',[]):
                writer.writerow([])
                writer.writerow([f"BẢNG {t['table_id']} / Trang {page['number']+1}"])
                for row in t['grid']:
                    writer.writerow([safe_value(str(v)) for v in row])
    return destination


def result_to_excel(result: dict, destination: Path):
    destination.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook()
    ws=wb.active;ws.title='Thông tin'
    ws.append(['HAK SMART OCR AI - KẾT QUẢ TRÍCH XUẤT'])
    ws.append(['Tên file',result.get('filename','')])
    ws.append(['Mã tài liệu',result.get('document_id','')])
    ws.append(['Loại tài liệu',result.get('document_type','')])
    ws.append(['Trạng thái',result.get('status','')])
    ws.append([])
    ws.append(['Trường dữ liệu','Giá trị OCR','Giá trị hiện tại','Đã xác nhận','Trang'])
    for f in result.get('fields',[]):
        ws.append([safe_value(f['label']),safe_value(f.get('raw','')),
                   safe_value(str(f.get('value',''))), 'Có' if f.get('verified') else 'Chưa', f['page']+1])
    ws.column_dimensions['A'].width=34;ws.column_dimensions['B'].width=38
    ws.column_dimensions['C'].width=38;ws.column_dimensions['D'].width=20
    ws.column_dimensions['E'].width=12
    count=0
    candidate_rows=[]
    for page in result.get('pages',[]):
        for tbl in page.get('tables',[]):
            count+=1
            out=wb.create_sheet(f'Bang_{page["number"]+1}_{count}'[:31])
            out.append([f'BẢNG TRANG {page["number"]+1} - {tbl["table_id"]}'])
            out.append(['Các ô còn cần đối chiếu với tài liệu gốc trước khi sử dụng nghiệp vụ'])
            out.append([])
            for row in tbl['grid']:
                out.append([safe_value(str(cell)) for cell in row])
            for suggestion in tbl.get('cell_suggestions',[]):
                ri,ci=suggestion['row'],suggestion['col']
                candidate_rows.append((tbl['table_id'],page['number']+1,ri+1,ci+1,
                                       safe_value(str(suggestion['value'])),
                                       suggestion.get('method',''),
                                       'Đã duyệt' if tbl.get('verified') else 'CHƯA DUYỆT'))
                # Data starts in row 4. Yellow cells represent guesses,
                # not verified payroll entries.
                if not tbl.get('verified') and ri < len(tbl['grid']) and ci < len(tbl['grid'][ri]):
                    out.cell(ri+4,ci+1).fill=PatternFill('solid',fgColor='FFF1BF')
                    out.cell(ri+4,ci+1).comment=None
            for col in range(1,min(50,max((len(x) for x in tbl['grid']),default=1))+1):
                out.column_dimensions[get_column_letter(col)].width=17
            out.freeze_panes='C4'
            out.sheet_view.showGridLines=True
            for row in out.iter_rows(min_row=4):
                for c in row:
                    c.alignment=Alignment(vertical='center',wrap_text=True)
                    c.border=Border(bottom=Side(style='hair',color='D5DDE5'),right=Side(style='hair',color='D5DDE5'))
    if candidate_rows:
        suggestions_ws=wb.create_sheet('Goi_y_OCR')
        suggestions_ws.append(['GỢI Ý OCR TRONG CÁC Ô — PHẢI ĐỐI CHIẾU ẢNH GỐC'])
        suggestions_ws.append(['Các giá trị do OCR/CV đề xuất, không phải dữ liệu chấm công đã được duyệt.'])
        suggestions_ws.append([])
        suggestions_ws.append(['Mã bảng','Trang','Dòng','Cột','Giá trị đề xuất','Phương pháp','Trạng thái'])
        for row in candidate_rows:suggestions_ws.append(row)
        suggestions_ws.freeze_panes='E5'
        suggestions_ws.auto_filter.ref=f'A4:G{suggestions_ws.max_row}'
        for col,width in {'A':27,'B':10,'C':10,'D':10,'E':24,'F':26,'G':22}.items():
            suggestions_ws.column_dimensions[col].width=width
    attendance=result.get('attendance_analysis',[])
    if attendance:
        attendance_sheet=wb.create_sheet('Cong_chuan_hoa')
        attendance_sheet.append(['NHẬT KÝ CÔNG OCR - KHÔNG PHẢI BẢNG LƯƠNG ĐƯỢC DUYỆT'])
        attendance_sheet.append(['Giữ nguyên ô trống; nguồn chưa duyệt không được đưa vào tính lương.'])
        attendance_sheet.append([])
        attendance_sheet.append(['Mã bảng','Trang','Kỳ','Nhân sự','Ngày','Chỉ tiêu',
                                 'Nội dung gốc','Giá trị số','Đã duyệt bảng','Ngày suy luận từ vị trí'])
        for table in attendance:
            for item in table['records']:
                attendance_sheet.append([safe_value(table['table_id']),table['page']+1,
                    table.get('period'),safe_value(item.get('person') or ''),item['day'],
                    safe_value(item['metric']),safe_value(str(item['raw'])),item['value'],
                    'Có' if table['verified'] else 'Chưa',
                    'Có - phải soát' if item.get('inferred_day') else 'Không'])
        attendance_sheet.auto_filter.ref=f'A4:J{max(attendance_sheet.max_row,4)}'
        attendance_sheet.freeze_panes='E5'
        for col,width in {'A':26,'B':8,'C':13,'D':29,'E':10,'F':16,'G':19,'H':12,'I':17,'J':26}.items():
            attendance_sheet.column_dimensions[col].width=width
        issues=wb.create_sheet('Canh_bao_OCR')
        issues.append(['Mã bảng','Trang','Trạng thái','Chi tiết cảnh báo'])
        for table in attendance:
            for issue in table['issues']:
                issues.append([safe_value(table['table_id']),table['page']+1,
                               'Chưa duyệt' if not table['verified'] else 'Cần kiểm tra',safe_value(issue)])
        issues.column_dimensions['D'].width=95
    summary=audit_attendance(result)
    if summary:
        audit=wb.create_sheet('Doi_chieu_cong')
        audit.append(['ĐỐI CHIẾU OCR - KHÔNG DÙNG ĐỂ TÍNH LƯƠNG CHƯA DUYỆT'])
        audit.append(['Số tổng là số đọc được; ô trống KHÔNG được coi là số 0.'])
        audit.append([])
        audit.append(['Mã bảng','Trang','Kỳ','Nhân sự','Chỉ tiêu','Tổng đã đọc',
                      'Ô có số','Ô trống','Ô không hợp lệ','Tổng ô','Đã duyệt','Cảnh báo'])
        for item in summary:
            audit.append([safe_value(item['table_id']),item['page']+1,item['period'],
                          safe_value(item['person'] or ''),safe_value(item['metric']),
                          item['read_sum'],item['numeric_cells'],item['blank_cells'],
                          item['invalid_cells'],item['expected_cells'],
                          'Có' if item['verified'] else 'Chưa',
                          safe_value('; '.join(item['warnings']))])
        audit.freeze_panes='F5'
        audit.auto_filter.ref=f'A4:L{audit.max_row}'
        for col,width in {'A':26,'B':8,'C':14,'D':28,'E':18,'F':15,
                          'G':12,'H':12,'I':16,'J':12,'K':12,'L':90}.items():
            audit.column_dimensions[col].width=width
    evidence=result.get('evidence_links',[])
    if evidence:
        ev=wb.create_sheet('Chung_tu_lien_quan')
        ev.append(['Quan hệ','Tên tệp liên quan','Trạng thái chứng từ','Ghi chú','Người tạo'])
        for item in evidence:
            ev.append([item['relationship'],safe_value(item['related_filename']),
                       item['related_status'],safe_value(item['note']),safe_value(item['reviewer'])])
        for col,width in {'A':20,'B':45,'C':24,'D':80,'E':20}.items():ev.column_dimensions[col].width=width
    for sh in wb.worksheets:
        sh['A1'].font=Font(name='Arial',size=14,bold=True,color='0E497B')
        sh.sheet_properties.pageSetUpPr.fitToPage=True
    wb.save(destination)
    return destination


def result_to_text(result: dict,destination: Path):
    destination.parent.mkdir(parents=True,exist_ok=True)
    chunks=[]
    for p in result.get('pages',[]):
        chunks.extend((f'======= TRANG {p["number"]+1} =======',p.get('text','')))
    destination.write_text('\n'.join(chunks),encoding='utf-8-sig')
    return destination


def export_result(result: dict, path: str | Path):
    path=Path(path)
    ext=path.suffix.lower()
    if ext=='.json':return result_to_json(result,path)
    if ext=='.xlsx':return result_to_excel(result,path)
    if ext=='.csv':return result_to_csv(result,path)
    if ext=='.txt':return result_to_text(result,path)
    raise ValueError('Định dạng xuất phải là JSON/XLSX/CSV/TXT')
