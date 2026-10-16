from hak_ocr.attendance_audit import audit_attendance
from hak_ocr.exporting import export_result
from openpyxl import load_workbook

def test_audit_separates_blanks_invalid_and_numeric():
    doc={'attendance_analysis':[{'table_id':'p0','page':0,'period':'09/2026',
        'verified':True,'records':[
        {'person':'A','metric':'cong','day':1,'value':1.0,'raw':'1','bbox':[0,0,.1,.1]},
        {'person':'A','metric':'cong','day':2,'value':None,'raw':'','bbox':None},
        {'person':'A','metric':'cong','day':3,'value':None,'raw':'I','bbox':None},
        {'person':'A','metric':'cong','day':4,'value':.5,'raw':'0.5','bbox':None}]}]}
    rows=audit_attendance(doc)
    assert rows[0]['read_sum']==1.5
    assert rows[0]['blank_cells']==1 and rows[0]['invalid_cells']==1
    assert rows[0]['payroll_ready'] is False

def test_export_contains_audit_not_payroll(tmp_path):
    doc={'filename':'sample','status':'NEEDS_REVIEW','fields':[],'pages':[],
         'attendance_analysis':[{'table_id':'p0','page':0,'period':None,'verified':False,
          'records':[{'person':'=CMD','metric':'cong','day':1,'value':1,'raw':'1'}],
          'issues':['Cần duyệt']}]}
    output=tmp_path/'x.xlsx'
    export_result(doc,output)
    wb=load_workbook(output)
    assert 'Doi_chieu_cong' in wb.sheetnames
    assert wb['Doi_chieu_cong']['D5'].value=="'=CMD"
    assert wb['Doi_chieu_cong']['K5'].value=='Chưa'
