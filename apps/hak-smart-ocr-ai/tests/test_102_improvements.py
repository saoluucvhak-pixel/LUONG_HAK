"""Regression tests for small-cell OCR provenance and stable human reviews."""
import json

import pytest
import numpy as np
from PIL import Image, ImageDraw
from hak_ocr.deep_cells import _single_stem_candidate, _read_number, refine_numeric_tables
from hak_ocr.models import Table,Page,Result
from hak_ocr.storage import Repository,table_layout_signature


def _make_table(xshift=0):
    boxes=[[[.10+xshift,.10,.40+xshift,.20],[.40+xshift,.10,.80+xshift,.20]],
           [[.10+xshift,.20,.40+xshift,.30],[.40+xshift,.20,.80+xshift,.30]]]
    return Table(0,[.10+xshift,.10,.80+xshift,.30],
                 [['Ngày','CÔNG'],['1','1']],boxes,table_id='page-0-grid-1')


def test_single_stem_visual_not_blank_or_two_strokes():
    blank=Image.new('L',(220,150),'white')
    assert not _single_stem_candidate(blank)
    stem=blank.copy()
    d=ImageDraw.Draw(stem)
    d.rectangle((109,42,115,112),fill='black')
    d.rectangle((106,42,115,46),fill='black')
    d.rectangle((106,109,119,112),fill='black')
    assert _single_stem_candidate(stem)
    d.rectangle((128,42,138,112),fill='black')
    assert not _single_stem_candidate(stem)
    assert _read_number(blank)==''


def test_table_source_layout_signature_changes_with_geometry():
    same=_make_table().to_dict()
    altered=_make_table(.04).to_dict()
    assert table_layout_signature(same)!=table_layout_signature(altered)
    modified=_make_table().to_dict()
    modified['grid'][1][1]='2'
    # OCR text alone is not a layout change.
    assert table_layout_signature(same)==table_layout_signature(modified)


def test_rerun_never_attaches_old_verified_table_to_changed_geometry(tmp_path):
    im=Image.new('RGB',(320,220),'white');source=tmp_path/'source.png';im.save(source)
    repo=Repository(tmp_path/'data')
    try:
        did,_=repo.add_file(source)
        d=repo.get(did)
        old=_make_table()
        repo.save_result(Result(did,d['sha256'],source.name,'bang_cham_cong',
                                [Page(0,320,220,'Tháng 09/2026',[],[old])]))
        corrected=[list(x) for x in old.grid]
        corrected[1][1]='2'
        repo.table_edit(did,old.table_id,0,corrected,True)
        before=repo.get_result(did)
        assert before['pages'][0]['tables'][0]['verified'] is True
        assert before['pages'][0]['tables'][0]['grid'][1][1]=='2'

        updated=_make_table(.04)
        repo.save_result(Result(did,d['sha256'],source.name,'bang_cham_cong',
                                [Page(0,320,220,'Tháng 09/2026',[],[updated])]))
        result=repo.get_result(did)
        now=result['pages'][0]['tables'][0]
        assert not now.get('verified')
        assert now['grid'][1][1]=='1'  # NOT the stale reviewed value
        assert 'đổi cấu trúc' in now['geometry_warning']
        assert old.table_id not in repo.verified_table_ids(did)
        with pytest.raises(ValueError,match='bảng chưa xác nhận'):
            repo.verify_document(did)
        # Corrected values remain in the history/database for manual recovery.
        previous=repo.db.execute('SELECT grid_json FROM table_edits WHERE document_id=?',(did,)).fetchone()[0]
        assert json.loads(previous)[1][1]=='2'
    finally:
        repo.close()


def test_invalid_month_day_rejected_without_zero_invention():
    from hak_ocr.attendance import analyse_attendance
    t=_make_table().to_dict()
    # A 30-day month with 31st table row must not be inferred as day 31.
    t['grid']=[['Ngày','CÔNG','CƠM']]+[[str(x),'1','1'] for x in range(1,31)]+[['116','30','29']]
    t['cell_boxes']=[]
    result={'pages':[{'number':0,'text':'Lịch làm Tháng 09/2026','tables':[t]}]}
    analysed=analyse_attendance(result)[0]
    assert len(analysed['records'])==60
    assert not any(rec['day']==31 for rec in analysed['records'])
    assert any('ngoài tháng' in issue for issue in analysed['issues'])


def test_header_recovery_requires_agreeing_readers(monkeypatch):
    from hak_ocr.deep_cells import refine_attendance_headers
    table=_make_table()
    table.grid[0][1]=''
    image=Image.new('RGB',(500,300),'white')
    monkeypatch.setattr('hak_ocr.deep_cells.pytesseract.image_to_string',
                        lambda *args,**kwargs:'CƠM')
    assert refine_attendance_headers(image,[table])==1
    assert table.grid[0][1]=='CƠM'
    assert table.cell_suggestions[0]['verified'] is False
    assert table.cell_suggestions[0]['method']=='header-two-pass-ocr'


def test_excel_export_marks_unreviewed_candidates(tmp_path):
    from openpyxl import load_workbook
    from hak_ocr.exporting import result_to_excel
    table=_make_table()
    table.cell_suggestions=[{'row':1,'col':1,'value':'1','verified':False,'method':'numeric-stem-cv'}]
    result={'filename':'sample.png','document_id':'demo','document_type':'bang_cham_cong',
            'status':'NEEDS_REVIEW','fields':[],
            'pages':[{'number':0,'tables':[table.to_dict()]}]}
    path=tmp_path/'sample.xlsx'
    result_to_excel(result,path)
    wb=load_workbook(path)
    assert 'Goi_y_OCR' in wb.sheetnames
    assert wb['Goi_y_OCR']['G5'].value=='CHƯA DUYỆT'
    assert wb['Bang_1_1']['B5'].fill.fgColor.rgb.endswith('FFF1BF')


def test_table_structural_edits_block_final_approval_when_bbox_lost(tmp_path):
    pic=tmp_path/'x.png';Image.new('RGB',(100,100),'white').save(pic)
    repo=Repository(tmp_path/'db')
    try:
        did,_=repo.add_file(pic);d=repo.get(did)
        t=_make_table()
        repo.save_result(Result(did,d['sha256'],'x.png','generic',
                                [Page(0,100,100,'',[],[t])]))
        wider=[row+[''] for row in t.grid]
        repo.table_edit(did,t.table_id,0,wider,True)
        result=repo.get_result(did)
        assert not result['pages'][0]['tables'][0]['verified']
        assert 'ánh xạ' in result['pages'][0]['tables'][0]['geometry_warning']
        with pytest.raises(ValueError,match='bảng chưa xác nhận'):
            repo.verify_document(did)
    finally:
        repo.close()
